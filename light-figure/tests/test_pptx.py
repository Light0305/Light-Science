"""v1.1 PPT 文件层全链（FIG-D02）：pptx manifest 契约、python-pptx
构建/反读对账、后端分流、powerpoint_verified 如实标注、原生性门
PPTX 分支与 round_noop 门。"""

import pytest

from light_figure.audit import run_audit, set_review_disposition
from light_figure.delivery import generate_delivery
from light_figure.gates import advance
from light_figure.pptxbuild import (
    PptxError, build_pptx, extract_and_reconcile_pptx,
    install_pptx_manifest, validate_pptx_manifest,
)
from light_figure.project import load_package, save_contract
from light_figure.provenance import update_provenance_row
from light_figure.util import dump_json, load_structured, sha256_file, today
from conftest import (
    FIG_DATA, FIG_METHOD, make_reference_png, method_manifest,
    schematic_contract,
)


def _pptx_manifest(root):
    manifest = method_manifest()
    manifest["version"] = "1.1"
    manifest["backend"] = "powerpoint_file"
    manifest["arrow_plan"] = [{
        "arrow_plan_id": "AP-1", "from": "input-frame",
        "to": "predictor", "fact": "参考图中输入框指向预测框的箭头",
    }]
    manifest["shapes"][2]["connect"]["arrow_plan_id"] = "AP-1"
    manifest["reference"]["sha256"] = sha256_file(
        root / "schematic" / "reference" / f"{FIG_METHOD}.png")
    return manifest


@pytest.fixture()
def pptx_pkg(filled_pkg):
    """方法图后端切到 powerpoint 并走完文件层全链（produced 门前）；
    数据线同 produced_pkg 口径生产，保证整包可推进。"""
    from light_figure.produce import produce_data_figure
    from light_figure.prompt import build_prompt
    from light_figure.reference import accept_reference, register_reference
    from light_figure.style import build_style
    from conftest import write_data_script, write_source_data

    root = filled_pkg["root"]
    contract = schematic_contract()
    contract["schematic_backend"] = "powerpoint"
    save_contract(root, contract)
    result = advance(root)
    assert result["advanced"], result["gaps"]
    bundle = build_style(root)
    write_source_data(root)
    write_data_script(root, bundle["style_build_hash"])
    produce_data_figure(root, FIG_DATA)
    build_prompt(root, FIG_METHOD)
    ref_src = make_reference_png(filled_pkg["tmp"] / "reference.png")
    register_reference(root, FIG_METHOD, ref_src,
                       model_name="GPT Image", model_version="1")
    accept_reference(root, FIG_METHOD, by="author", checks={
        "labels_verbatim": True, "structures_factual": True,
        "formulas_values_checked": True})
    manifest_file = filled_pkg["tmp"] / "pptx-manifest.json"
    dump_json(manifest_file, _pptx_manifest(root))
    install_pptx_manifest(root, FIG_METHOD, manifest_file)
    build_pptx(root, FIG_METHOD)
    outcome = extract_and_reconcile_pptx(root, FIG_METHOD)
    assert outcome["ok"], outcome["diffs"]
    update_provenance_row(root, FIG_METHOD, {
        "label_check": {"done": True, "by": "author", "at": today()},
    }, by="test")
    return filled_pkg


def test_pptx_full_chain_to_delivered(pptx_pkg):
    root = pptx_pkg["root"]
    result = advance(root)  # contracted → produced（PPTX 分流门）
    assert result["advanced"], result["gaps"]
    run_audit(root)
    record = load_structured(root / "audit" / f"{FIG_METHOD}.audit.json")
    gates = record["gates"]
    assert gates["native"]["verdict"] == "PASS"
    assert gates["contract"]["verdict"] == "PASS"
    assert gates["policy_block"]["verdict"] == "PASS"
    assert record["artifact"].endswith(".pptx")
    for fig_id in (FIG_DATA, FIG_METHOD):
        rec = load_structured(root / "audit" / f"{fig_id}.audit.json")
        for name, gate in rec["gates"].items():
            if gate.get("applicable") and gate["verdict"] == \
                    "REVIEW_REQUIRED":
                set_review_disposition(root, fig_id, name,
                                       "已人工复核：容差内可接受")
    result = advance(root)  # produced → audited
    assert result["advanced"], result["gaps"]
    out = generate_delivery(root)
    assert out["numbering_frozen"] is True
    doc = load_structured(root / "exports" / "delivery.yaml")
    row = next(r for r in doc["figures"] if r["fig_id"] == FIG_METHOD)
    assert row["files"][0]["path"].endswith(".pptx")
    assert row["files"][0]["role"] == "powerpoint_source"
    assert row["powerpoint_verified"] is False
    assert "未过 PowerPoint 实机验证" in row["powerpoint_note"]
    corr = row["correspondence"]
    assert corr["content_hashes"].get("extracted_shapes")
    assert any(s["role"] == "redraw_manifest"
               for s in corr["source_assets"])
    result = advance(root)  # audited → delivered
    assert result["advanced"], result["gaps"]


def test_pptx_extract_ids_pinned(pptx_pkg):
    from light_figure.pptxbuild import extract_pptx_shapes
    root = pptx_pkg["root"]
    shapes = extract_pptx_shapes(
        root / "schematic" / "redraw" / f"{FIG_METHOD}.pptx")
    by_name = {s["shape_id"]: s for s in shapes}
    assert by_name["input-frame"]["user_semantic_id"] == "input-frame"
    assert "arrow-AP-1" in by_name
    assert by_name["arrow-AP-1"]["arrow"]["from"] == "input-frame"
    assert by_name["arrow-AP-1"]["arrow"]["to"] == "predictor"


def test_pptx_native_gate_blocks_full_slide_picture(pptx_pkg):
    from pptx import Presentation
    from pptx.util import Inches

    root = pptx_pkg["root"]
    path = root / "schematic" / "redraw" / f"{FIG_METHOD}.pptx"
    prs = Presentation(str(path))
    slide = prs.slides[0]
    slide.shapes.add_picture(
        str(root / "schematic" / "reference" / f"{FIG_METHOD}.png"),
        0, 0, prs.slide_width, prs.slide_height)
    prs.save(str(path))
    run_audit(root, fig_id=FIG_METHOD)
    record = load_structured(root / "audit" / f"{FIG_METHOD}.audit.json")
    native = record["gates"]["native"]
    assert native["verdict"] == "FIX_BEFORE_DELIVERY"
    assert any(f["code"] in ("native.undeclared_bitmap",
                             "native.full_page_bitmap")
               for f in native["findings"])


def test_pptx_verified_false_claim_blocked(pptx_pkg):
    root = pptx_pkg["root"]
    mpath = root / "schematic" / "redraw" / \
        f"{FIG_METHOD}.pptx-manifest.json"
    manifest = load_structured(mpath)
    manifest["build"]["powerpoint_verified"] = True
    dump_json(mpath, manifest)
    run_audit(root, fig_id=FIG_METHOD)
    record = load_structured(root / "audit" / f"{FIG_METHOD}.audit.json")
    contract = record["gates"]["contract"]
    assert contract["verdict"] == "FIX_BEFORE_DELIVERY"
    assert any(f["code"] == "contract.powerpoint_verified_false_claim"
               for f in contract["findings"])


def test_pptx_round_noop_gate(pptx_pkg):
    root = pptx_pkg["root"]
    with pytest.raises(PptxError) as excinfo:
        build_pptx(root, FIG_METHOD)
    assert any("round_noop" in g for g in excinfo.value.gaps)


def test_pptx_manifest_validation_catches_bad_geometry(pkg):
    root = pkg["root"]
    ref = root / "schematic" / "reference" / f"{FIG_METHOD}.png"
    ref.parent.mkdir(parents=True, exist_ok=True)
    make_reference_png(ref)
    manifest = _pptx_manifest(root)
    manifest["shapes"][0]["box_local"] = [0.5, 0.5, 0.9, 0.9]
    gaps = validate_pptx_manifest(manifest)
    assert any("越界" in g for g in gaps)
    manifest = _pptx_manifest(root)
    manifest["shapes"][2]["connect"]["arrow_plan_id"] = "AP-404"
    gaps = validate_pptx_manifest(manifest)
    assert any("AP-404" in g for g in gaps)
