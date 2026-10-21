"""P3 渲染层：申报副本 → DOCX + 回读核验（段落/编号/关键字段
集合比对）、过程词拒渲染、缺申报副本点名。"""

import pytest
from docx import Document

from light_patent import draft, render
from light_patent.project import PatentError, case_path, load_structured

from conftest import drive_to_drafted


def _drafted_and_stripped(case_root, tmp_path):
    drive_to_drafted(case_root, tmp_path)
    draft.strip_filing(case_root)
    return case_root


def test_render_three_docs_verified(case_root, tmp_path):
    _drafted_and_stripped(case_root, tmp_path)
    result = render.render_docx(case_root)
    assert result["files"] == 3
    assert result["all_verified"] is True
    filing = case_path(case_root, "filing_dir")
    for name in ("权利要求书.docx", "说明书.docx", "说明书摘要.docx"):
        assert (filing / name).exists()
    report = load_structured(filing / "render_report.yaml")
    assert report["kind"] == "light-patent/render-report"
    assert all(f["verify"]["ok"] for f in report["files"])
    assert all(f["docx_sha256"] for f in report["files"])


def test_rendered_claims_match_source_numbers(case_root, tmp_path):
    _drafted_and_stripped(case_root, tmp_path)
    render.render_docx(case_root)
    doc = Document(str(case_path(case_root, "filing_dir") / "权利要求书.docx"))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    claim_lines = [t for t in texts if t[:2] in ("1.", "2.", "3.", "4.")]
    assert len(claim_lines) == 4
    assert claim_lines[0].startswith("1. 一种家禽异常行为检测方法")
    # 页眉为文件名称、标题不带工作副本字样
    header_text = doc.sections[0].header.paragraphs[0].text
    assert header_text == "权利要求书"
    assert all("工作副本" not in t for t in texts)


def test_rendered_spec_sections_present(case_root, tmp_path):
    _drafted_and_stripped(case_root, tmp_path)
    render.render_docx(case_root)
    doc = Document(str(case_path(case_root, "filing_dir") / "说明书.docx"))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    for section in ("技术领域", "背景技术", "发明内容", "附图说明", "具体实施方式"):
        assert section in texts
    assert texts[0] == "说明书"


def test_render_refuses_unclean_source(case_root, tmp_path):
    drive_to_drafted(case_root, tmp_path)
    draft.strip_filing(case_root)
    # 人工往申报副本里塞过程词：渲染层复扫必须拒（LP-04 不许旁路）
    spec = case_path(case_root, "filing_dir") / "specification.md"
    spec.write_text(spec.read_text(encoding="utf-8") + "\n查新过程记录\n",
                    encoding="utf-8")
    with pytest.raises(PatentError) as ei:
        render.render_docx(case_root)
    assert "拒渲染" in str(ei.value)


def test_render_requires_filing_first(case_root, tmp_path):
    drive_to_drafted(case_root, tmp_path)  # 未 strip
    with pytest.raises(PatentError) as ei:
        render.render_docx(case_root)
    assert "strip" in str(ei.value)


def test_verify_docx_detects_tamper(case_root, tmp_path):
    _drafted_and_stripped(case_root, tmp_path)
    render.render_docx(case_root)
    filing = case_path(case_root, "filing_dir")
    paras = render.parse_filing_md(
        (filing / "claims.md").read_text(encoding="utf-8"))
    verdict = render.verify_docx(filing / "权利要求书.docx", paras)
    assert verdict["ok"] and verdict["claims"] == 4
    # 篡改：删一段后核验必须咬合
    doc = Document(str(filing / "权利要求书.docx"))
    p = doc.paragraphs[-1]
    p._element.getparent().remove(p._element)
    doc.save(str(filing / "权利要求书.docx"))
    verdict2 = render.verify_docx(filing / "权利要求书.docx", paras)
    assert not verdict2["ok"]
