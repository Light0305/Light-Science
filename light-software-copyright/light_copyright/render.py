"""SC4 源程序鉴别材料（确定性渲染）+ SC5 文档鉴别材料（草稿段）。

SC4 纪律（Q5 代码轨）：分页由构造保证——有效行（非空源行按显示列折行）
逐行装页，每页恰 50 行（末页除外）；总量超 60 页取前后各连续 30 页，
断口与省略量显形登记。每一材料行在索引里回指（文件， 原始行号）
（HC-1 可回溯）。渲染后再实测复核（构造值 vs 实测值双记）。

SC5 纪律（Q2 折中 + HC-3）：工具只起草稿段。每段带源码证据指针
（模块/符号由 AST 机械抽取，指针必须真存在）；无证据的功能段不生成。
段落落 draft/segments/*.md + manifest，确认走 case 的回执机制；
本模块只产草稿，绝不自称定版。
"""

from __future__ import annotations

import ast
import json
from pathlib import Path

from . import schema
from .case import CaseError, log_tool_use
from .select import load_selection
from .snapshot import load_snapshot
from .util import dump_json, load_structured, sha256_text, wrap_line

MATERIALS_DIR = "materials"
DRAFT_DIR = "draft"
PAGE_HEADER_PREFIX = "[页眉]"


def header_line(facts: dict) -> str:
    sw = facts.get("software") or {}
    return f"{sw.get('full_name', '')} {sw.get('version', '')}".strip()


# ------------------------------------------------------------------ SC4
def build_code_lines(root, facts) -> tuple:
    """把入选文件按快照顺序展成材料行流：（文本， 文件， 原始行号）。"""
    root = Path(root)
    snap = load_snapshot(root)
    sel = load_selection(root)
    source_dir = Path(snap["source_dir"])
    meta_by_path = {f["path"]: f for f in snap["files"]}
    stream = []
    for entry in sel["included"]:
        rel = entry["path"]
        if rel not in meta_by_path:
            raise CaseError([f"入选文件不在冻结快照内：{rel}"])
        text = (source_dir / rel).read_text(encoding="utf-8")
        for lineno, raw in enumerate(text.splitlines(), start=1):
            if not raw.strip():
                continue  # 空行不计有效行（每页 50 行口径为非空有效行）
            for piece in wrap_line(raw.rstrip(), schema.WRAP_COLUMNS):
                stream.append({"text": piece, "file": rel,
                               "line": lineno})
    return stream, snap


def paginate(stream: list, per_page: int) -> list:
    return [stream[i:i + per_page]
            for i in range(0, len(stream), per_page)]


def front_back(pages: list) -> tuple:
    """前后各 30 页规则：60 页及以下全交；超限取首尾，断口信息返回。"""
    n = len(pages)
    if n <= schema.FULL_SUBMIT_MAX_PAGES:
        return pages, None
    head = pages[:schema.FRONT_BACK_PAGES]
    tail = pages[-schema.FRONT_BACK_PAGES:]
    omitted = {
        "omitted_pages": n - 2 * schema.FRONT_BACK_PAGES,
        "omitted_from_page": schema.FRONT_BACK_PAGES + 1,
        "omitted_to_page": n - schema.FRONT_BACK_PAGES,
        "omitted_lines": sum(len(p) for p in
                             pages[schema.FRONT_BACK_PAGES:
                                   n - schema.FRONT_BACK_PAGES]),
    }
    return head + tail, omitted


def render_code_material(root, facts) -> dict:
    root = Path(root)
    stream, snap = build_code_lines(root, facts)
    pages, omitted = front_back(paginate(stream, schema.CODE_LINES_PER_PAGE))
    hdr = header_line(facts)
    out_lines, index = [], []
    total_pages = len(pages)
    full_pages = (len(stream) + schema.CODE_LINES_PER_PAGE - 1) // \
        schema.CODE_LINES_PER_PAGE
    for pi, page in enumerate(pages, start=1):
        if omitted is None or pi <= schema.FRONT_BACK_PAGES:
            real_no = pi
        else:
            real_no = full_pages - total_pages + pi
        out_lines.append(f"{PAGE_HEADER_PREFIX}{hdr}  第 {real_no} 页")
        for row in page:
            out_lines.append(row["text"])
            index.append({"material_page": real_no, "file": row["file"],
                          "source_line": row["line"]})
        out_lines.append("")
    mdir = root / MATERIALS_DIR
    mdir.mkdir(parents=True, exist_ok=True)
    text_path = mdir / "源程序鉴别材料.txt"
    text_path.write_text("\n".join(out_lines), encoding="utf-8")
    with (mdir / "源程序鉴别材料.index.jsonl").open("w",
                                                  encoding="utf-8") as fh:
        for row in index:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")
    manifest = {
        "kind": "code",
        "file": text_path.name,
        "constructed_pages": total_pages,
        "full_page_count": full_pages,
        "lines_per_page": schema.CODE_LINES_PER_PAGE,
        "material_lines": len(index),
        "omitted": omitted,
        "header": hdr,
        "page_basis": "constructed（构造保证）+ 实测复核",
    }
    log_tool_use(root, "tool", "render", "code-material",
                 f"{total_pages} 页（全量 {full_pages} 页）")
    return manifest


def measure_code_material(root, manifest: dict) -> dict:
    """实测兜底：重读材料文本复核页数与每页行数（构造值的独立核对）。"""
    mdir = Path(root) / MATERIALS_DIR
    text = (mdir / manifest["file"]).read_text(encoding="utf-8")
    pages, cur = [], None
    for line in text.splitlines():
        if line.startswith(PAGE_HEADER_PREFIX):
            if cur is not None:
                pages.append(cur)
            cur = 0
        elif line.strip() and cur is not None:
            cur += 1
    if cur is not None:
        pages.append(cur)
    bad = [i + 1 for i, c in enumerate(pages[:-1])
           if c != schema.CODE_LINES_PER_PAGE]
    return {"measured_pages": len(pages), "page_line_counts": pages,
            "bad_pages": bad}


# ------------------------------------------------------------------ SC5
def extract_symbols(source_dir: Path, rel: str) -> dict:
    """AST 机械抽取模块符号（证据指针的来源；抽不到如实记空）。"""
    path = source_dir / rel
    info = {"module_doc": "", "functions": [], "classes": []}
    if path.suffix != ".py":
        return info
    try:
        tree = ast.parse(path.read_text(encoding="utf-8"))
    except (SyntaxError, OSError, UnicodeDecodeError):
        return info
    info["module_doc"] = (ast.get_docstring(tree) or "").strip()
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            info["functions"].append(node.name)
        elif isinstance(node, ast.ClassDef):
            info["classes"].append(node.name)
    return info


def draft_doc_segments(root, facts) -> dict:
    """按 Q3 分流生成文档草稿段（每段带证据指针），落 draft/segments。"""
    root = Path(root)
    kind = (facts.get("software") or {}).get("kind")
    doc_type = schema.DOC_TYPE_BY_KIND.get(kind or "")
    if not doc_type:
        raise CaseError([f"software.kind={kind!r} 无法分流文档形态（Q3）"])
    snap = load_snapshot(root)
    sel = load_selection(root)
    source_dir = Path(snap["source_dir"])
    seg_dir = root / DRAFT_DIR / "segments"
    seg_dir.mkdir(parents=True, exist_ok=True)
    for old in seg_dir.glob("SEG-*.md"):
        old.unlink()
    sw = facts["software"]
    segments = []

    def add(section, text, evidence):
        seg_id = f"SEG-{len(segments) + 1:02d}"
        path = seg_dir / f"{seg_id}.md"
        if not text.endswith("\n"):
            text += "\n"
        path.write_text(text, encoding="utf-8")
        segments.append({
            "seg_id": seg_id, "section": section,
            "file": str(path.relative_to(root)), "doc_type": doc_type,
            "evidence": evidence, "sha256": sha256_text(text),
        })

    add("封面与概述",
        f"# {sw['full_name']} {doc_type}\n\n"
        f"软件全称：{sw['full_name']}\n版本号：{sw['version']}\n"
        f"著作权人：{(facts.get('owner') or {}).get('name', '')}\n\n"
        f"本{doc_type}对应软件 {sw['full_name']}（版本 {sw['version']}），"
        "内容由登记事实单源与冻结源码快照派生，功能陈述逐段回指源码证据。",
        [{"file": "(facts)", "symbol": "software.full_name"}])
    for entry in sel["included"]:
        rel = entry["path"]
        info = extract_symbols(source_dir, rel)
        symbols = info["functions"] + info["classes"]
        if not symbols and not info["module_doc"]:
            continue  # 无可指证据的文件不生成成功能段（HC-3）
        desc = info["module_doc"] or "（无模块说明，功能据符号清单陈述）"
        sym_text = "、".join(symbols) if symbols else "（无顶层函数/类）"
        add(f"模块：{rel}",
            f"## 模块 {rel}\n\n{desc}\n\n"
            f"本模块对外提供：{sym_text}。"
            "上述符号由源码 AST 机械抽取，与证据指针逐一对应；"
            "段落内不陈述证据之外的功能。",
            [{"file": rel, "symbol": s} for s in symbols] or
            [{"file": rel, "symbol": "(module)"}])
    env = facts.get("environment") or {}
    if env:
        env_lines = "\n".join(f"- {k}：{v}" for k, v in env.items() if v)
        add("运行与支撑环境",
            f"## 运行与支撑环境\n\n{env_lines}\n\n"
            "本段由登记事实 environment 区派生，与申请表环境字段同源。",
            [{"file": "(facts)", "symbol": "environment"}])
    manifest = {"doc_type": doc_type, "segments": segments}
    dump_json(root / DRAFT_DIR / "segments-manifest.json", manifest)
    log_tool_use(root, "tool", "draft", "doc-segments",
                 f"{doc_type} 草稿 {len(segments)} 段")
    return manifest


def load_segments(root) -> dict:
    p = Path(root) / DRAFT_DIR / "segments-manifest.json"
    if not p.exists():
        raise CaseError(["尚无文档草稿段（先 draft-doc）"])
    return load_structured(p)


def segment_text(root, seg: dict) -> str:
    return (Path(root) / seg["file"]).read_text(encoding="utf-8")


def render_doc_material(root, facts) -> dict:
    """把确认后的段落拼成文档鉴别材料：30 行/页构造分页 + DOCX 轨。"""
    root = Path(root)
    man = load_segments(root)
    hdr = header_line(facts)
    full = []
    for seg in man["segments"]:
        full.extend(segment_text(root, seg).splitlines())
    stream = []
    for line in full:
        if not line.strip():
            continue
        stream.extend(wrap_line(line, 60))
    pages, omitted = front_back(paginate(
        [{"text": t} for t in stream], schema.DOC_LINES_PER_PAGE))
    out = []
    for pi, page in enumerate(pages, start=1):
        out.append(f"{PAGE_HEADER_PREFIX}{hdr}  第 {pi} 页")
        out.extend(r["text"] for r in page)
        out.append("")
    mdir = root / MATERIALS_DIR
    mdir.mkdir(parents=True, exist_ok=True)
    doc_type = man["doc_type"]
    name = f"文档鉴别材料-{doc_type}.txt"
    (mdir / name).write_text("\n".join(out), encoding="utf-8")
    manifest = {"kind": "doc", "doc_type": doc_type, "file": name,
                "constructed_pages": len(pages), "omitted": omitted,
                "lines_per_page": schema.DOC_LINES_PER_PAGE,
                "header": hdr, "page_basis": "constructed（文本模型）"}
    _render_docx(root, facts, man, manifest)
    log_tool_use(root, "tool", "render", "doc-material",
                 f"{doc_type} {len(pages)} 页")
    return manifest


def _render_docx(root, facts, segments_manifest, doc_manifest) -> None:
    """Q5 文档轨：DOCX 渲染 + 目录页码按文本模型回填（如实标注口径）。"""
    try:
        from docx import Document
    except ImportError:
        doc_manifest["docx"] = "unavailable（python-docx 未安装，已登记）"
        return
    doc = Document()
    doc.add_heading(f"{header_line(facts)} "
                    f"{segments_manifest['doc_type']}", level=0)
    doc.add_heading("目录（页码按文本模型分页回填，非办公套件真分页）",
                    level=1)
    per = schema.DOC_LINES_PER_PAGE
    line_no = 0
    toc, bodies = [], []
    for seg in segments_manifest["segments"]:
        text = segment_text(root, seg)
        toc.append((seg["section"], line_no // per + 2))
        bodies.append(text)
        line_no += sum(1 for ln in text.splitlines() if ln.strip())
    for section, page in toc:
        doc.add_paragraph(f"{section} …… {page}")
    for text in bodies:
        for ln in text.splitlines():
            doc.add_paragraph(ln)
    name = f"文档鉴别材料-{segments_manifest['doc_type']}.docx"
    doc.save(str(Path(root) / MATERIALS_DIR / name))
    doc_manifest["docx"] = name
    doc_manifest["toc_basis"] = "文本模型 30 行/页回填；真分页以办公套件打开为准"
